from __future__ import annotations

import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from docx import Document
from sklearn.cluster import KMeans
from sklearn.metrics import silhouette_score
from sklearn.preprocessing import StandardScaler

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw"
PROCESSED_DIR = ROOT / "data" / "processed"
FIG_DIR = ROOT / "outputs" / "figures"
TABLE_DIR = ROOT / "outputs" / "tables"


def _normalise_text(series: pd.Series) -> pd.Series:
    return series.astype(str).str.strip().str.replace(r"\s+", " ", regex=True)


def load_raw_dataset() -> pd.DataFrame:
    frames = []
    for path in sorted(RAW_DIR.glob("*.csv")):
        df = pd.read_csv(path)
        frames.append(df)
    if not frames:
        raise FileNotFoundError(f"No CSV files found in {RAW_DIR}")
    data = pd.concat(frames, ignore_index=True)
    return data


def clean_dataset(df: pd.DataFrame) -> pd.DataFrame:
    data = df.copy()
    required_cols = [
        "Complaint ID",
        "Category",
        "Sub Category",
        "Grievance Date",
        "Ward Name",
        "Grievance Status",
        "Staff Remarks",
        "Staff Name",
    ]
    missing = [c for c in required_cols if c not in data.columns]
    if missing:
        raise ValueError(f"Missing columns: {missing}")

    data["Complaint ID"] = pd.to_numeric(data["Complaint ID"], errors="coerce")
    data = data.dropna(subset=["Complaint ID"]).drop_duplicates(subset=["Complaint ID"], keep="first")

    for col in ["Category", "Sub Category", "Ward Name", "Grievance Status", "Staff Name"]:
        data[col] = _normalise_text(data[col]).replace({"nan": "Unknown", "None": "Unknown", "": "Unknown"})

    data["Category"] = data["Category"].str.title()
    data["Sub Category"] = data["Sub Category"].str.title()
    data["Ward Name"] = data["Ward Name"].str.title()
    data["Grievance Status"] = data["Grievance Status"].str.title().replace({"Nan": "Unknown", "Na": "Unknown"})
    data["Staff Name"] = data["Staff Name"].str.title()

    data["Grievance Date"] = pd.to_datetime(data["Grievance Date"], errors="coerce")
    data = data.dropna(subset=["Grievance Date"]).copy()

    data["Year"] = data["Grievance Date"].dt.year
    data["Month"] = data["Grievance Date"].dt.month_name()
    data["Quarter"] = data["Grievance Date"].dt.to_period("Q").astype(str)
    data["DayOfWeek"] = data["Grievance Date"].dt.day_name()
    data["Grievance Status"] = data["Grievance Status"].fillna("Unknown")

    data = data.sort_values("Grievance Date").reset_index(drop=True)
    return data


def compute_kpis(df: pd.DataFrame) -> dict:
    yearly = df["Year"].value_counts().sort_index()
    category_counts = df["Category"].value_counts()
    ward_counts = df["Ward Name"].value_counts()
    status_counts = df["Grievance Status"].value_counts()

    kpis = {
        "total_grievances": int(len(df)),
        "unique_wards": int(df["Ward Name"].nunique()),
        "unique_categories": int(df["Category"].nunique()),
        "unique_subcategories": int(df["Sub Category"].nunique()),
        "top_category": category_counts.index[0],
        "top_category_count": int(category_counts.iloc[0]),
        "top_ward": ward_counts.index[0],
        "top_ward_count": int(ward_counts.iloc[0]),
        "yearly_counts": yearly.to_dict(),
        "status_counts": status_counts.to_dict(),
        "category_top_10": category_counts.head(10).to_dict(),
        "ward_top_10": ward_counts.head(10).to_dict(),
    }
    return kpis


def generate_visuals(df: pd.DataFrame) -> dict:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    summary = {}

    sns.set_style("whitegrid")
    plt.figure(figsize=(10, 5))
    yearly = df["Year"].value_counts().sort_index()
    yearly.plot(kind="bar", color="#4f46e5")
    plt.title("Grievance Volume by Year")
    plt.xlabel("Year")
    plt.ylabel("Complaints")
    plt.tight_layout()
    path = FIG_DIR / "yearly_trend.png"
    plt.savefig(path, dpi=180)
    plt.close()
    summary["yearly_trend"] = str(path)

    plt.figure(figsize=(10, 6))
    cat = df["Category"].value_counts().head(10)
    cat.plot(kind="barh", color="#14b8a6")
    plt.title("Top 10 Complaint Categories")
    plt.xlabel("Complaints")
    plt.ylabel("Category")
    plt.tight_layout()
    path = FIG_DIR / "category_distribution.png"
    plt.savefig(path, dpi=180)
    plt.close()
    summary["category_distribution"] = str(path)

    plt.figure(figsize=(10, 6))
    ward = df["Ward Name"].value_counts().head(10)
    ward.plot(kind="barh", color="#f59e0b")
    plt.title("Top 10 Highest Demand Wards")
    plt.xlabel("Complaints")
    plt.ylabel("Ward")
    plt.tight_layout()
    path = FIG_DIR / "ward_demand.png"
    plt.savefig(path, dpi=180)
    plt.close()
    summary["ward_demand"] = str(path)

    status = df["Grievance Status"].value_counts().head(8)
    plt.figure(figsize=(8, 8))
    plt.pie(status.values, labels=status.index, autopct="%1.1f%%", startangle=90)
    plt.title("Grievance Status Distribution")
    plt.tight_layout()
    path = FIG_DIR / "status_distribution.png"
    plt.savefig(path, dpi=180)
    plt.close()
    summary["status_distribution"] = str(path)

    return summary


def build_ml_segments(df: pd.DataFrame) -> dict:
    top_cats = df["Category"].value_counts().head(8).index.tolist()
    ward_agg = df.groupby("Ward Name").agg(
        total_grievances=("Complaint ID", "count"),
        category_diversity=("Category", "nunique"),
        status_closed=("Grievance Status", lambda s: (s == "Closed").mean()),
        status_rejected=("Grievance Status", lambda s: (s == "Rejected").mean()),
    )

    for cat in top_cats:
        ward_agg[f"share_{cat}"] = df.groupby("Ward Name")["Category"].apply(
            lambda x: (x == cat).mean()
        )

    features = ward_agg[[c for c in ward_agg.columns if c != "total_grievances"]]
    scaler = StandardScaler()
    X = scaler.fit_transform(features.fillna(0))

    best_k = 2
    best_score = -1
    for k in range(2, min(6, len(ward_agg)) + 1):
        model = KMeans(n_clusters=k, random_state=42, n_init=10)
        labels = model.fit_predict(X)
        score = silhouette_score(X, labels)
        if score > best_score:
            best_score = score
            best_k = k

    final_model = KMeans(n_clusters=best_k, random_state=42, n_init=10)
    cluster_labels = final_model.fit_predict(X)
    ward_agg["cluster"] = cluster_labels

    cluster_summary = (
        ward_agg.groupby("cluster")
        .agg(
            wards=("total_grievances", "count"),
            avg_grievances=("total_grievances", "mean"),
            avg_category_diversity=("category_diversity", "mean"),
            pct_closed=("status_closed", "mean"),
            pct_rejected=("status_rejected", "mean"),
        )
        .sort_values("avg_grievances", ascending=False)
    )

    return {
        "best_k": best_k,
        "silhouette_score": round(float(best_score), 4),
        "cluster_summary": cluster_summary,
        "ward_clusters": ward_agg.reset_index()[["Ward Name", "cluster", "total_grievances"]],
    }


def write_summary_json(kpis: dict, figures: dict, segmentation: dict) -> None:
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    summary = {
        "project": "CivicPulse",
        "dataset_period": "2020-2024",
        "kpis": kpis,
        "figures": figures,
        "ml": {
            "best_k": segmentation["best_k"],
            "silhouette_score": segmentation["silhouette_score"],
            "cluster_summary": segmentation["cluster_summary"].to_dict(),
        },
    }
    with (PROCESSED_DIR / "civicpulse_summary.json").open("w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2, default=str)


def write_report_doc(kpis: dict, segmentation: dict) -> None:
    doc = Document()
    doc.add_heading("CivicPulse — BBMP Grievance Intelligence", 0)
    doc.add_paragraph("Ward-wise Civic Demand, Grievance Trends & Municipal Decision Analytics")
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        f"This project analyzed {kpis['total_grievances']:,} BBMP grievance records from 2020–2024. "
        f"The dominant complaint category was {kpis['top_category']} with {kpis['top_category_count']:,} records. "
        f"The most complaint-heavy ward was {kpis['top_ward']} with {kpis['top_ward_count']:,} cases."
    )

    doc.add_heading("Key KPI Snapshot", level=1)
    for label, value in {
        "Total grievances": f"{kpis['total_grievances']:,}",
        "Unique wards": f"{kpis['unique_wards']}",
        "Unique categories": f"{kpis['unique_categories']}",
        "Unique subcategories": f"{kpis['unique_subcategories']}",
        "Top category": kpis["top_category"],
        "Top ward": kpis["top_ward"],
        "ML clusters": str(segmentation["best_k"]),
        "Silhouette score": str(segmentation["silhouette_score"]),
    }.items():
        doc.add_paragraph(f"{label}: {value}")

    doc.add_heading("Analytical Interpretation", level=1)
    doc.add_paragraph(
        "Electrical complaints dominate the grievance profile, followed by solid waste and road maintenance complaints. "
        "This indicates recurring municipal service issues with both urban infrastructure and civic maintenance implications."
    )
    doc.add_paragraph(
        "Ward demand is unevenly distributed. High-demand wards such as Jnanabharathi Ward and Rajarajeshwari Nagar show sustained complaint concentration, suggesting the need for ward-level monitoring and prioritization."
    )

    doc.save(ROOT / "DeepakDubey_CivicPulse_ProjectReport.docx")


def run_project() -> dict:
    raw = load_raw_dataset()
    cleaned = clean_dataset(raw)
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    cleaned.to_csv(PROCESSED_DIR / "bbmp_grievances_clean.csv", index=False)

    kpis = compute_kpis(cleaned)
    figures = generate_visuals(cleaned)
    segmentation = build_ml_segments(cleaned)
    write_summary_json(kpis, figures, segmentation)
    write_report_doc(kpis, segmentation)

    return {"data": cleaned, "kpis": kpis, "figures": figures, "ml": segmentation}


if __name__ == "__main__":
    run_project()
